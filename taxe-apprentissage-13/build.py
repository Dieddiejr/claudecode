"""Génère l'Excel et les données de la carte à partir de organismes_ta_13.csv.

Usage : python3 build.py
Sorties : organismes_ta_13.xlsx, organismes_ta_13.json, carte_ta_13.html
"""
import csv
import json
import zlib
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent

# Centroïdes approximatifs (commune ou arrondissement de Marseille).
# Les points de la carte sont placés à ce niveau, pas à l'adresse exacte.
CENTROIDES = {
    "Marseille": (43.2965, 5.3810),
    "Marseille 1er": (43.3000, 5.3840),
    "Marseille 2e": (43.3120, 5.3660),
    "Marseille 3e": (43.3120, 5.3800),
    "Marseille 4e": (43.3060, 5.4000),
    "Marseille 5e": (43.2930, 5.3980),
    "Marseille 6e": (43.2870, 5.3800),
    "Marseille 7e": (43.2830, 5.3600),
    "Marseille 8e": (43.2500, 5.3800),
    "Marseille 9e": (43.2450, 5.4300),
    "Marseille 10e": (43.2770, 5.4270),
    "Marseille 11e": (43.2900, 5.4800),
    "Marseille 12e": (43.3070, 5.4400),
    "Marseille 13e": (43.3500, 5.4300),
    "Marseille 14e": (43.3450, 5.3950),
    "Marseille 15e": (43.3600, 5.3650),
    "Marseille 16e": (43.3600, 5.3200),
    "Aix-en-Provence": (43.5290, 5.4470),
    "Vitrolles": (43.4600, 5.2480),
    "Salon-de-Provence": (43.6400, 5.0970),
    "Arles": (43.6770, 4.6310),
    "Martigues": (43.4050, 5.0480),
    "Gémenos": (43.2950, 5.6280),
    "Aubagne": (43.2920, 5.5700),
    "Miramas": (43.5820, 5.0020),
    "Barbentane": (43.9000, 4.7490),
    "La Roque-d'Anthéron": (43.7160, 5.3100),
    "Eyragues": (43.8410, 4.8420),
    "Puyloubier": (43.5250, 5.6760),
    "Saint-Martin-de-Crau": (43.6400, 4.8130),
    "Lambesc": (43.6540, 5.2620),
    "Rousset": (43.4830, 5.6200),
    "Istres": (43.5140, 4.9880),
    "Bouches-du-Rhône": None,
}

FIABILITE = {
    "A": "A – Collecte TA / habilitation SOLTéA affichée par l'organisme",
    "B": "B – Statut éligible (L6241-5) ; inscription SOLTéA 2026 à vérifier",
    "C": "C – Éligibilité ou statut incertain, à confirmer",
}

EXCLUS = [
    ("CCI Aix-Marseille-Provence (écoles et CFA consulaires)", "Établissement public administratif (consulaire)", "https://taxe-apprentissage.cciamp.com/"),
    ("Chambre de métiers et de l'artisanat PACA", "Établissement public administratif (consulaire)", "https://www.cmar-paca.fr/article/soutenez-lartisanat-avec-la-taxe-dapprentissage"),
    ("Aix-Marseille Université et ses composantes (IUT, Polytech, facultés, INSPE, IMPGT…)", "Établissement public", "https://polytech.univ-amu.fr/en/companies/apprenticeship-tax"),
    ("Centrale Méditerranée, Arts et Métiers Aix, École de l'air et de l'espace", "Établissements publics", "https://www.safecluster.com/actualites/etablissements-habilites-a-percevoir-le-solde-de-la-taxe-dapprentissage-dans-le-reseau-du-pole-safe/"),
    ("ENSA-Marseille, Beaux-Arts de Marseille (ESADMM), Cnam PACA", "Établissements publics", "https://www.cnam-paca.fr/alternance-entreprises"),
    ("Lycées publics de l'académie (ex. La Viste, Antonin Artaud…)", "Établissements publics", "https://www.site.ac-aix-marseille.fr/lyc-viste/spip/Campagne-Taxe-Apprentissage-2023.html"),
    ("Écoles privées à but lucratif (SAS/SA : groupes Ionis, Galileo, etc.)", "Non éligibles au solde : L6241-5 vise les organismes à but non lucratif", "https://entreprendre.service-public.gouv.fr/vosdroits/F22574"),
    ("CFA (part apprentissage)", "Financés par les 87 % via les OPCO ; le solde ne leur est versé qu'en équipements (L6241-4)", "https://www.centre-inffo.fr/site-droit-formation/modalites-de-gestion-de-la-part-de-13-de-la-taxe-dapprentissage"),
    ("La Bourguette (autisme)", "Association, mais ses établissements fléchés TA sont dans le Vaucluse", "https://www.bourguette-autisme.org/taxe-d-apprentissage/"),
]


def coordonnees(row):
    """Centroïde de la commune, décalé de façon stable pour éviter les superpositions."""
    base = CENTROIDES.get(row["commune"])
    if base is None:
        return None, None
    h = zlib.crc32(row["id"].encode())
    dlat = ((h & 0xFF) / 255 - 0.5) * 0.010
    dlon = (((h >> 8) & 0xFF) / 255 - 0.5) * 0.014
    return round(base[0] + dlat, 5), round(base[1] + dlon, 5)


def lire():
    with open(HERE / "organismes_ta_13.csv", encoding="utf-8") as f:
        rows = list(csv.DictReader(f, delimiter=";"))
    for r in rows:
        r["lat"], r["lon"] = coordonnees(r)
    return rows


def ecrire_xlsx(rows):
    wb = Workbook()
    entete = Font(bold=True, color="FFFFFF")
    fond = PatternFill("solid", fgColor="1F4E5A")
    couleurs = {"A": "D7EFD9", "B": "FFF2CC", "C": "F8D7D3"}

    ws = wb.active
    ws.title = "Organismes"
    cols = ["id", "nom", "categorie", "base_legale", "gestionnaire_statut", "adresse", "cp",
            "commune", "uai", "siret", "fiabilite", "note", "source", "lat", "lon"]
    titres = ["N°", "Organisme", "Catégorie", "Base légale", "Gestionnaire / statut", "Adresse", "CP",
              "Commune", "UAI", "SIRET", "Fiabilité", "Remarque", "Source", "Lat. (approx.)", "Lon. (approx.)"]
    ws.append(titres)
    for c in ws[1]:
        c.font, c.fill = entete, fond
    for r in rows:
        ws.append([int(r["id"]) if k == "id" else r[k] for k in cols])
        ws.cell(ws.max_row, cols.index("fiabilite") + 1).fill = PatternFill("solid", fgColor=couleurs[r["fiabilite"]])
        src = ws.cell(ws.max_row, cols.index("source") + 1)
        src.hyperlink, src.style = r["source"], "Hyperlink"
    largeurs = [5, 48, 30, 18, 36, 34, 8, 20, 22, 17, 9, 60, 50, 12, 12]
    for i, w in enumerate(largeurs, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    s = wb.create_sheet("Synthèse")
    s.append(["Catégorie", "A", "B", "C", "Total"])
    for c in s[1]:
        c.font, c.fill = entete, fond
    par_cat = {}
    for r in rows:
        par_cat.setdefault(r["categorie"], Counter())[r["fiabilite"]] += 1
    for cat, cnt in sorted(par_cat.items(), key=lambda kv: -sum(kv[1].values())):
        s.append([cat, cnt["A"], cnt["B"], cnt["C"], sum(cnt.values())])
    tot = Counter(r["fiabilite"] for r in rows)
    s.append(["Total", tot["A"], tot["B"], tot["C"], len(rows)])
    s[s.max_row][0].font = Font(bold=True)
    s.append([])
    for k, v in FIABILITE.items():
        s.append([v])
    s.column_dimensions["A"].width = 70

    e = wb.create_sheet("Exclus (publics, lucratifs)")
    e.append(["Organisme", "Motif d'exclusion", "Source"])
    for c in e[1]:
        c.font, c.fill = entete, fond
    for row in EXCLUS:
        e.append(list(row))
    e.column_dimensions["A"].width = 70
    e.column_dimensions["B"].width = 80
    e.column_dimensions["C"].width = 60

    m = wb.create_sheet("Méthode")
    texte = [
        "Périmètre : organismes non publics (associations, fondations, établissements privés sous contrat ou non lucratifs) "
        "situés dans les Bouches-du-Rhône et susceptibles de percevoir le solde (13 %) de la taxe d'apprentissage via SOLTéA.",
        "Base légale : article L6241-5 du Code du travail ; listes préfectorale et régionale PACA 2026 (arrêté du 26/05/2026, liste complémentaire au 24/07/2026).",
        "Limite : la liste officielle (XLSX de la préfecture PACA, plateforme SOLTéA) n'a pas pu être téléchargée depuis l'environnement de travail. "
        "Le recensement est reconstitué à partir des pages publiques des organismes et d'annuaires (Onisep, L'Etudiant, enseignement-prive.info).",
        "Pour un recensement exhaustif : télécharger la liste préfectorale PACA 2026 et lancer filtrer_liste_officielle.py.",
        "Coordonnées : centroïde de la commune ou de l'arrondissement (pas l'adresse exacte).",
        "Calendrier 2026 : fléchage SOLTéA du 26/05 au 21/08 puis du 03/09 au 21/10/2026 (clos à la date du recensement, 23/10/2026).",
    ]
    for t in texte:
        m.append([t])
        m.cell(m.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    m.column_dimensions["A"].width = 130

    wb.save(HERE / "organismes_ta_13.xlsx")


def main():
    rows = lire()
    ecrire_xlsx(rows)
    data = {"organismes": rows, "exclus": [dict(zip(("nom", "motif", "source"), e)) for e in EXCLUS]}
    (HERE / "organismes_ta_13.json").write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    page = (HERE / "carte_template.html").read_text(encoding="utf-8")
    page = page.replace("/*DATA*/", json.dumps(data, ensure_ascii=False).replace("</", "<\\/"))
    (HERE / "carte_ta_13.html").write_text(page, encoding="utf-8")
    print(len(rows), "organismes ;", Counter(r["fiabilite"] for r in rows))


if __name__ == "__main__":
    main()
